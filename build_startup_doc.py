from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "STARTUP_GUIDE.md"
OUTPUT = ROOT / "IITB_Integrated_Security_System_Startup_Guide.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_border(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def remove_style_border(style):
    properties = style._element.get_or_add_pPr()
    border = properties.find(qn("w:pBdr"))
    if border is not None:
        properties.remove(border)


def add_inline(paragraph, text):
    pieces = re.split(r"(`[^`]+`)", text)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, "Consolas", 9.5)
        else:
            run = paragraph.add_run(piece)
            set_run_font(run)


def add_code_block(doc, code):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F4F7")
    properties.append(shading)
    run = paragraph.add_run(code.rstrip())
    set_run_font(run, "Consolas", 8.7)
    return paragraph


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.autofit = False
    widths = [Inches(0.85), Inches(3.65), Inches(0.85)]
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.width = widths[index] if index < len(widths) else Inches(1.0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value.strip())
        set_run_font(run, "Aptos", 9.5, True, "FFFFFF")
    for row_index, row in enumerate(rows[2:], start=1):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            cell.width = widths[col_index] if col_index < len(widths) else Inches(1.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "F3F6FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value.strip())
            set_run_font(run, "Aptos", 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.5)
    section.footer_distance = Inches(0.25)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (("Title", 23, "000000"), ("Heading 1", 15, "000000"), ("Heading 2", 12, "000000")):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        remove_style_border(style)

    title = next((line[2:].strip() for line in lines if line.startswith("# ")), "IITB Integrated Security System Startup Guide")
    title_paragraph = doc.add_paragraph(style="Title")
    title_paragraph.paragraph_format.space_after = Pt(4)
    remove_paragraph_border(title_paragraph)
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, "Aptos Display", 23, True)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    add_inline(intro, "This manual explains how to clone, prepare, start, and verify the IITB Integrated Security System using seven separate PowerShell terminals on Windows.")

    index = 0
    first_heading_skipped = False
    in_code = False
    code_lines = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("# ") and not first_heading_skipped:
            first_heading_skipped = True
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and lines[index + 1].startswith("| ---"):
            table_rows = []
            while index < len(lines) and lines[index].startswith("|"):
                values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                table_rows.append(values)
                index += 1
            add_table(doc, table_rows)
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(11)
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run(line[3:].strip())
            set_run_font(run, "Aptos", 15, True)
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(line[4:].strip())
            set_run_font(run, "Aptos", 12, True)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, line[2:].strip())
        elif line.strip():
            paragraph = doc.add_paragraph()
            add_inline(paragraph, line.strip())
        index += 1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("IITB Integrated Security System | Startup Guide")
    set_run_font(footer_run, "Aptos", 8, False, "666666")
    doc.core_properties.title = title
    doc.core_properties.subject = "Manual startup instructions for the seven project services"
    doc.core_properties.author = "IITB Integrated Security System"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
